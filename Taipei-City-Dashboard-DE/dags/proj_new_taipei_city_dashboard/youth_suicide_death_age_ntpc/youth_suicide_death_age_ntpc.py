from airflow import DAG
from operators.common_pipeline import CommonDag


SOURCE_URL = (
    "https://www.mohw.gov.tw/dl-83420-8ac74f07-080f-4b94-8453-ce2a8ee99f8b.html"
)
UA = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
    "AppleWebKit/537.36 Chrome/131.0 Safari/537.36"
)
NTPC_COUNTY_CODE = "31"
SUICIDE_CAUSE_CODE = "131"

CONTRACT_COLUMNS = [
    "indicator_id", "period_start", "period_end", "period_type",
    "age_lower", "age_upper", "age_band_raw", "gender", "area_code",
    "area_level", "breakdown", "value", "unit", "value_type", "data_time",
]

ROLE_ALIASES = {
    "year": ("year", "年度", "年度別"),
    "county": ("county", "縣市", "縣市代碼", "鄉鎮市區代碼"),
    "cause": ("cause", "死因", "死亡原因"),
    "sex": ("sex", "性別"),
    "age_code": ("age_code", "年齡代碼"),
    "value": ("N", "死亡數", "count", "value"),
}

KNOWN_BAD_YEARS = {}
SCALE_ANOMALY_FACTOR = 3.0


def _normalise(value):
    import unicodedata

    return unicodedata.normalize("NFKC", str(value)).strip().replace(" ", "")


def _text(value, field):
    value = "" if value is None else str(value).strip()
    if not value or value.lower() in {"nan", "none", "null"}:
        raise ValueError(f"自殺死因來源欄位 {field} 為空")
    return value


def _number(value, field):
    raw = _text(value, field)
    try:
        number = float(raw.replace(",", ""))
    except ValueError as exc:
        raise ValueError(f"自殺死因來源欄位 {field} 無法轉數值：{value!r}") from exc
    if number < 0 or number != number:
        raise ValueError(f"自殺死因來源欄位 {field} 出現無效數值：{value!r}")
    return number


def _column_indexes(headers):
    normalised = {_normalise(value): index for index, value in enumerate(headers)}
    indexes = {}
    for role, aliases in ROLE_ALIASES.items():
        matches = [normalised[_normalise(alias)] for alias in aliases
                   if _normalise(alias) in normalised]
        if len(matches) != 1:
            raise ValueError(
                f"自殺死因來源找不到唯一的 {role} 欄位；實際欄位={headers}、"
                f"別名={aliases}、命中={matches}"
            )
        indexes[role] = matches[0]
    return indexes


def _codebook(zfile):
    """讀取 ZIP 內官方對照表，避免自行猜 age_code 或 cause code。"""
    import io

    import openpyxl

    workbook_name = next(
        (name for name in zfile.namelist() if name.lower().endswith(".xlsx")),
        None,
    )
    if workbook_name is None:
        raise RuntimeError("自殺死因 ZIP 找不到官方 XLSX 對照表。")
    workbook = openpyxl.load_workbook(
        io.BytesIO(zfile.read(workbook_name)), data_only=True, read_only=True
    )
    sheet_names = {_normalise(name): name for name in workbook.sheetnames}
    if "age_code" not in sheet_names or "cause" not in sheet_names:
        raise RuntimeError(
            f"自殺死因 ZIP 對照表缺少 age_code/cause 工作表：{workbook.sheetnames}"
        )
    age_sheet = workbook[sheet_names["age_code"]]
    age_map = {}
    for row in age_sheet.iter_rows(min_row=2, values_only=True):
        if row[1] is None or row[2] is None:
            continue
        age_map[str(row[1]).strip().zfill(2)] = _text(row[2], "age_code label")
    cause_sheet = workbook[sheet_names["cause"]]
    cause_map = {}
    for row in cause_sheet.iter_rows(min_row=2, values_only=True):
        if row[0] is None or row[1] is None:
            continue
        cause_map[str(int(row[0])).zfill(3)] = _text(row[1], "cause label")
    if SUICIDE_CAUSE_CODE not in cause_map:
        raise RuntimeError("官方 cause 對照表找不到自殺代碼 131。")
    if "自殺" not in cause_map[SUICIDE_CAUSE_CODE] and "自我傷害" not in cause_map[SUICIDE_CAUSE_CODE]:
        raise RuntimeError(
            f"cause 131 對照表語意不是自殺：{cause_map[SUICIDE_CAUSE_CODE]!r}"
        )
    return age_map, cause_map


def _parse_age(age_code, age_label):
    import re

    code = _text(age_code, "age_code").zfill(2)
    label = _text(age_label, "age_label")
    if code in {"01", "02"}:
        # 新生兒／嬰兒不是可精確表達成「整數歲」的區間；不把它猜成 0 歲。
        return None, None, "non_year_age_band"
    if code == "99":
        if "不詳" not in label and "未知" not in label:
            raise ValueError(f"age_code 99 的官方標籤不似不詳：{label!r}")
        return None, None, "unknown_age"
    if code in {"03", "04", "05", "06"}:
        # 官方對照表使用「一歲」至「四歲」中文單齡標籤。
        age = int(code) - 2
        return age, age, "single_age"
    match = re.fullmatch(r"(\d{1,3})-(\d{1,3})歲", label)
    if match:
        lower, upper = int(match.group(1)), int(match.group(2))
        if lower > upper or int(code) != 7 + (lower - 5) // 5:
            raise ValueError(f"自殺死因 age_code 與五齡標籤不一致：{code!r}, {label!r}")
        return lower, upper, "age_band"
    match = re.fullmatch(r"(\d{1,3})歲以上", label)
    if match and code == "26":
        return int(match.group(1)), None, "age_band"
    raise ValueError(f"無法解析自殺死因來源年齡標籤：{code!r}, {label!r}")


def _gender(code):
    value = _text(code, "sex")
    mapping = {"1": "male", "2": "female"}
    if value not in mapping:
        raise ValueError(f"無法解析自殺死因來源性別代碼：{code!r}")
    return mapping[value]


def fetch_records():
    import csv
    import io
    import re
    import zipfile

    import requests

    response = requests.get(SOURCE_URL, headers={"User-Agent": UA}, timeout=120)
    response.raise_for_status()
    with zipfile.ZipFile(io.BytesIO(response.content)) as zfile:
        age_map, cause_map = _codebook(zfile)
        records = []
        csv_names = sorted(
            name for name in zfile.namelist()
            if re.fullmatch(r"opendata\d+\.csv", name, re.IGNORECASE)
        )
        if not csv_names:
            raise RuntimeError("自殺死因 ZIP 找不到 opendata*.csv。")
        for name in csv_names:
            roc_year = int(re.search(r"(\d+)", name).group(1))
            with zfile.open(name) as stream:
                reader = csv.reader(io.TextIOWrapper(stream, encoding="utf-8-sig"))
                try:
                    headers = next(reader)
                except StopIteration as exc:
                    raise RuntimeError(f"自殺死因檔案為空：{name}") from exc
                indexes = _column_indexes(headers)
                selected = 0
                for row_number, row in enumerate(reader, start=2):
                    if len(row) != len(headers):
                        raise ValueError(
                            f"自殺死因 {name} 第 {row_number} 列欄數不符："
                            f"{len(row)} != {len(headers)}"
                        )
                    county = _text(row[indexes["county"]], "county").zfill(2)
                    cause = _text(row[indexes["cause"]], "cause").zfill(3)
                    if county != NTPC_COUNTY_CODE or cause != SUICIDE_CAUSE_CODE:
                        continue
                    age_code = _text(row[indexes["age_code"]], "age_code").zfill(2)
                    if age_code not in age_map:
                        raise ValueError(f"官方 age_code 對照表找不到：{age_code!r}")
                    records.append(
                        {
                            "roc_year": roc_year,
                            "county": county,
                            "cause": cause,
                            "cause_label": cause_map[cause],
                            "sex": _text(row[indexes["sex"]], "sex"),
                            "age_code": age_code,
                            "age_label": age_map[age_code],
                            "value": _number(row[indexes["value"]], "value"),
                        }
                    )
                    selected += 1
            print(f"suicide source ROC {roc_year}: {selected} New Taipei rows")
    if not records:
        raise RuntimeError("自殺死因來源沒有新北市 cause=131 資料。")
    return records


def _year_totals(records):
    totals = {}
    for record in records:
        year = int(_text(record.get("roc_year"), "roc_year"))
        totals[year] = totals.get(year, 0.0) + _number(record.get("value"), "value")
    return totals


def drop_scope_anomalies(records):
    totals = _year_totals(records)
    if not totals:
        raise RuntimeError("自殺死因來源找不到年度總數，無法做尺度檢核。")
    kept = {year: value for year, value in totals.items() if year not in KNOWN_BAD_YEARS}
    if len(kept) >= 3:
        ordered = sorted(kept.values())
        middle = len(ordered) // 2
        median = ordered[middle] if len(ordered) % 2 else (ordered[middle - 1] + ordered[middle]) / 2
        for year, value in sorted(kept.items()):
            if median and (value > median * SCALE_ANOMALY_FACTOR
                           or value * SCALE_ANOMALY_FACTOR < median):
                raise ValueError(
                    f"民國 {year} 年新北市自殺死亡數 {value:.0f}，與其餘年度中位數 "
                    f"{median:.0f} 相差超過 {SCALE_ANOMALY_FACTOR} 倍；"
                    "請確認來源範圍或口徑後再更新 KNOWN_BAD_YEARS。"
                )
    else:
        print("suicide scale guard: fewer than 3 years; comparison not applicable")
    for year, reason in sorted(KNOWN_BAD_YEARS.items()):
        if year in totals:
            print(f"suicide scale guard: exclude ROC {year}: {reason}")
    return [
        record for record in records
        if int(_text(record.get("roc_year"), "roc_year")) not in KNOWN_BAD_YEARS
    ]


def _reconcile(records, data):
    expected = {}
    for record in records:
        year = int(_text(record.get("roc_year"), "roc_year"))
        expected[year] = expected.get(year, 0.0) + _number(record.get("value"), "value")
    if len(data) != len(records):
        raise ValueError(f"自殺死因輸出列數不一致：輸入 {len(records)}、輸出 {len(data)}")
    for roc_year, expected_total in sorted(expected.items()):
        start = f"{roc_year + 1911}-01-01"
        emitted = float(data.loc[data["period_start"] == start, "value"].sum())
        if abs(emitted - expected_total) > 0.5:
            raise ValueError(
                f"民國 {roc_year} 年自殺死亡數對帳失敗："
                f"輸入 {expected_total}、輸出 {emitted}"
            )
        print(f"ROC {roc_year}: input total {int(expected_total)}, output total {int(emitted)}, reconciled")


def transform_records(records, data_time=None):
    """DAG 與 verify 共用的五齡死因轉換核心。"""
    import json
    from datetime import datetime, timedelta, timezone

    import pandas as pd

    if not records:
        raise RuntimeError("自殺死因來源轉換前沒有資料列。")
    rows = []
    for record in records:
        roc_year = int(_text(record.get("roc_year"), "roc_year"))
        age_lower, age_upper, age_scope = _parse_age(
            record.get("age_code"), record.get("age_label")
        )
        gender = _gender(record.get("sex"))
        age_label = _text(record.get("age_label"), "age_label")
        rows.append(
            {
                "indicator_id": "suicide_death_count",
                "period_start": f"{roc_year + 1911}-01-01",
                "period_end": f"{roc_year + 1911}-12-31",
                "period_type": "year",
                "age_lower": age_lower,
                "age_upper": age_upper,
                "age_band_raw": age_label,
                "gender": gender,
                "area_code": "65000",
                "area_level": "city",
                "breakdown": json.dumps(
                    {
                        "cause_code": SUICIDE_CAUSE_CODE,
                        "cause_label": _text(record.get("cause_label"), "cause_label"),
                        "county_code": NTPC_COUNTY_CODE,
                        "age_code": _text(record.get("age_code"), "age_code"),
                        "age_scope": age_scope,
                    },
                    ensure_ascii=False,
                    sort_keys=True,
                ),
                "value": _number(record.get("value"), "value"),
                "unit": "人",
                "value_type": "count",
                "data_time": data_time,
            }
        )
    data = pd.DataFrame(rows, columns=CONTRACT_COLUMNS)
    data["age_lower"] = pd.array(data["age_lower"], dtype="Int16")
    data["age_upper"] = pd.array(data["age_upper"], dtype="Int16")
    if data_time is None:
        data["data_time"] = datetime.now(timezone(timedelta(hours=8))).isoformat()
    _reconcile(records, data)
    return data


def _transfer(**kwargs):
    from sqlalchemy import create_engine
    from utils.get_time import get_tpe_now_time_str
    from utils.load_stage import (
        save_dataframe_to_postgresql,
        update_lasttime_in_data_to_dataset_info,
    )

    dag_infos = kwargs.get("dag_infos")
    dag_id = dag_infos.get("dag_id")
    data = transform_records(
        drop_scope_anomalies(fetch_records()),
        data_time=get_tpe_now_time_str(is_with_tz=True),
    )
    print(f"ready_data shape =========== {data.shape}")
    engine = create_engine(kwargs.get("ready_data_db_uri"))
    save_dataframe_to_postgresql(
        engine,
        data=data,
        load_behavior=dag_infos.get("load_behavior"),
        default_table=dag_infos.get("ready_data_default_table"),
    )
    update_lasttime_in_data_to_dataset_info(engine, dag_id, data["data_time"].max())


dag = CommonDag(
    proj_folder="proj_new_taipei_city_dashboard",
    dag_folder="youth_suicide_death_age_ntpc",
)
dag.create_dag(etl_func=_transfer)
